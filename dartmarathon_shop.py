from __future__ import annotations

import asyncio
import csv
import io
import json
import os
import re
import secrets
import time
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any, Awaitable, Callable, Dict
from urllib.parse import urlencode

import aiohttp
from aiohttp import web


class ShopError(Exception):
    pass


class ShopManager:
    """Public Foltershop + PayPal checkout + admin task/article management."""

    def __init__(
        self,
        prefix: str,
        base_dir: Path,
        template_dir: Path,
    ) -> None:
        self.prefix = prefix.rstrip("/")
        self.template_dir = template_dir
        main_state_path = os.getenv("DARTMARATHON_STATE_FILE", "").strip()
        if main_state_path:
            default_shop_file = str(
                Path(main_state_path).with_name("dartmarathon_shop_state.json")
            )
        else:
            default_shop_file = str(base_dir / "dartmarathon_shop_state.json")

        self.state_file = Path(
            os.getenv(
                "DARTMARATHON_SHOP_STATE_FILE",
                default_shop_file,
            )
        )
        self.public_base_url = os.getenv(
            "DARTMARATHON_PUBLIC_BASE_URL",
            "",
        ).strip().rstrip("/")
        self.paypal_mode = os.getenv(
            "PAYPAL_MODE",
            "sandbox",
        ).strip().lower()
        if self.paypal_mode not in {"sandbox", "live"}:
            self.paypal_mode = "sandbox"
        self.paypal_client_id = os.getenv("PAYPAL_CLIENT_ID", "").strip()
        self.paypal_client_secret = os.getenv("PAYPAL_CLIENT_SECRET", "").strip()
        self.paypal_webhook_id = os.getenv("PAYPAL_WEBHOOK_ID", "").strip()
        self.reservation_seconds = max(
            180,
            min(
                1800,
                int(os.getenv("DARTMARATHON_SHOP_RESERVATION_SECONDS", "600") or 600),
            ),
        )

        self._lock = asyncio.Lock()
        self._state = self._load_state()
        self._paypal_token = ""
        self._paypal_token_expires_at = 0.0
        self._broadcast_state: Callable[[], Awaitable[None]] | None = None
        self._broadcast_payload: Callable[[Dict[str, Any]], Awaitable[None]] | None = None
        self._get_event_state: Callable[[], Dict[str, Any]] | None = None
        self._apply_event_action: Callable[
            [str, int, int],
            Awaitable[Dict[str, Any]],
        ] | None = None

    @property
    def paypal_api_base(self) -> str:
        if self.paypal_mode == "live":
            return "https://api-m.paypal.com"
        return "https://api-m.sandbox.paypal.com"

    @property
    def paypal_configured(self) -> bool:
        return bool(self.paypal_client_id and self.paypal_client_secret)

    @property
    def webhook_configured(self) -> bool:
        return bool(self.paypal_configured and self.paypal_webhook_id)

    def set_callbacks(
        self,
        broadcast_state: Callable[[], Awaitable[None]],
        broadcast_payload: Callable[[Dict[str, Any]], Awaitable[None]],
        get_event_state: Callable[[], Dict[str, Any]] | None = None,
        apply_event_action: Callable[
            [str, int, int],
            Awaitable[Dict[str, Any]],
        ] | None = None,
    ) -> None:
        self._broadcast_state = broadcast_state
        self._broadcast_payload = broadcast_payload
        self._get_event_state = get_event_state
        self._apply_event_action = apply_event_action

    def _default_state(self) -> Dict[str, Any]:
        return {
            "version": 3,
            "test_mode": False,
            "items": [],
            "orders": [],
            "next_item_id": 1,
            "next_order_id": 1,
        }

    def _load_state(self) -> Dict[str, Any]:
        if not self.state_file.exists():
            return self._default_state()
        try:
            raw = json.loads(self.state_file.read_text(encoding="utf-8"))
        except Exception as exc:
            print(f"[SHOP] State konnte nicht geladen werden: {exc}")
            return self._default_state()
        return self._sanitize_state(raw)

    def _sanitize_state(self, raw: Any) -> Dict[str, Any]:
        state = self._default_state()
        if not isinstance(raw, dict):
            return state

        state["test_mode"] = bool(raw.get("test_mode", False))

        items: list[Dict[str, Any]] = []
        for item in raw.get("items", []):
            cleaned = self._sanitize_item(item)
            if cleaned:
                items.append(cleaned)
        items.sort(key=lambda x: (int(x.get("sort_order", x["id"] * 10)), x["id"]))
        state["items"] = items

        orders: list[Dict[str, Any]] = []
        for order in raw.get("orders", []):
            cleaned = self._sanitize_order(order)
            if cleaned:
                orders.append(cleaned)
        orders.sort(key=lambda x: x["id"])
        state["orders"] = orders[-2000:]

        max_item = max((item["id"] for item in items), default=0)
        max_order = max((order["id"] for order in orders), default=0)
        try:
            state["next_item_id"] = max(max_item + 1, int(raw.get("next_item_id", max_item + 1)))
        except (TypeError, ValueError):
            state["next_item_id"] = max_item + 1
        try:
            state["next_order_id"] = max(max_order + 1, int(raw.get("next_order_id", max_order + 1)))
        except (TypeError, ValueError):
            state["next_order_id"] = max_order + 1
        return state

    def _sanitize_item(self, raw: Any) -> Dict[str, Any] | None:
        if not isinstance(raw, dict):
            return None
        try:
            item_id = int(raw.get("id", 0))
        except (TypeError, ValueError):
            return None
        if item_id <= 0:
            return None
        target = str(raw.get("target", "general"))
        if target not in {"tzmarty", "korsar", "both", "general", "on_site"}:
            target = "general"

        action_type = str(raw.get("action_type", "")).strip().lower()
        if action_type not in {"", "pause_minus"}:
            action_type = ""

        try:
            action_value = max(
                0,
                min(24 * 3600, int(raw.get("action_value", 0) or 0)),
            )
        except (TypeError, ValueError):
            action_value = 0

        try:
            sort_order = int(raw.get("sort_order", item_id * 10) or item_id * 10)
        except (TypeError, ValueError):
            sort_order = item_id * 10

        image_url = str(raw.get("image_url", "")).strip()[:600]
        if image_url and not image_url.lower().startswith(("http://", "https://")):
            image_url = ""

        return {
            "id": item_id,
            "name": str(raw.get("name", "Artikel"))[:80],
            "description": str(raw.get("description", ""))[:320],
            "price_cents": max(50, int(raw.get("price_cents", 100) or 100)),
            "max_quantity": max(0, min(999, int(raw.get("max_quantity", 0) or 0))),
            "cooldown_seconds": max(0, min(24 * 3600, int(raw.get("cooldown_seconds", 0) or 0))),
            "target": target,
            "stream_text": str(raw.get("stream_text", ""))[:220],
            "icon": str(raw.get("icon", "🎯"))[:12] or "🎯",
            "image_url": image_url,
            "sort_order": sort_order,
            "action_type": action_type,
            "action_value": action_value,
            "active": bool(raw.get("active", True)),
            "created_at": str(raw.get("created_at", "")),
            "updated_at": str(raw.get("updated_at", "")),
        }

    def _sanitize_order(self, raw: Any) -> Dict[str, Any] | None:
        if not isinstance(raw, dict):
            return None
        try:
            order_id = int(raw.get("id", 0))
        except (TypeError, ValueError):
            return None
        if order_id <= 0:
            return None
        kind = str(raw.get("kind", "donation"))
        if kind not in {"item", "donation"}:
            kind = "donation"
        status = str(raw.get("status", "created"))
        if status not in {"created", "paid", "cancelled", "expired", "failed"}:
            status = "failed"
        task_status = str(raw.get("task_status", "none"))
        if task_status not in {"none", "open", "in_progress", "done", "cancelled"}:
            task_status = "none"
        try:
            item_id = int(raw.get("item_id", 0) or 0)
        except (TypeError, ValueError):
            item_id = 0
        return {
            "id": order_id,
            "kind": kind,
            "item_id": item_id,
            "item_name": str(raw.get("item_name", ""))[:80],
            "target": str(raw.get("target", "general"))[:20],
            "stream_text": str(raw.get("stream_text", ""))[:220],
            "amount_cents": max(0, int(raw.get("amount_cents", 0) or 0)),
            "display_name": str(raw.get("display_name", ""))[:50],
            "message": str(raw.get("message", ""))[:180],
            "status": status,
            "task_status": task_status,
            "created_at": str(raw.get("created_at", "")),
            "created_ts": float(raw.get("created_ts", 0.0) or 0.0),
            "expires_at": float(raw.get("expires_at", 0.0) or 0.0),
            "paid_at": str(raw.get("paid_at", "")),
            "paid_ts": float(raw.get("paid_ts", 0.0) or 0.0),
            "paypal_order_id": str(raw.get("paypal_order_id", ""))[:100],
            "capture_id": str(raw.get("capture_id", ""))[:120],
            "is_test": bool(raw.get("is_test", False)),
            "action_type": (
                str(raw.get("action_type", "")).strip().lower()
                if str(raw.get("action_type", "")).strip().lower() in {"", "pause_minus"}
                else ""
            ),
            "action_value": max(0, int(raw.get("action_value", 0) or 0)),
            "action_status": (
                str(raw.get("action_status", "none"))
                if str(raw.get("action_status", "none")) in {"none", "pending", "applied"}
                else "none"
            ),
        }

    def _save_locked(self) -> None:
        try:
            self.state_file.parent.mkdir(parents=True, exist_ok=True)
            temp = self.state_file.with_suffix(self.state_file.suffix + ".tmp")
            temp.write_text(
                json.dumps(self._state, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            os.replace(temp, self.state_file)
        except Exception as exc:
            print(f"[SHOP] State konnte nicht gespeichert werden: {exc}")

    @staticmethod
    def _now_text() -> str:
        return datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    @staticmethod
    def _money(cents: int) -> str:
        return f"{max(0, cents) / 100:.2f}"

    @staticmethod
    def _parse_money_to_cents(
        value: Any,
        minimum_cents: int = 50,
        maximum_cents: int = 100000,
    ) -> int:
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            amount = float(value)
        else:
            text = (
                str(value or "")
                .strip()
                .replace("\xa0", "")
                .replace("€", "")
                .replace("EUR", "")
                .replace("eur", "")
                .replace(" ", "")
            )
            if not text:
                raise ShopError("Ungültiger Betrag.")

            if "," in text and "." in text:
                # deutsches 1.234,56 oder internationales 1,234.56
                if text.rfind(",") > text.rfind("."):
                    text = text.replace(".", "").replace(",", ".")
                else:
                    text = text.replace(",", "")
            elif "," in text:
                text = text.replace(",", ".")

            try:
                amount = float(text)
            except ValueError as exc:
                raise ShopError("Ungültiger Betrag.") from exc

        cents = int(round(amount * 100))
        if cents < minimum_cents or cents > maximum_cents:
            raise ShopError(
                f"Betrag muss zwischen {minimum_cents / 100:.2f} € "
                f"und {maximum_cents / 100:.2f} € liegen."
            )
        return cents

    def _cleanup_expired_locked(self) -> bool:
        now = time.time()
        changed = False
        for order in self._state["orders"]:
            if (
                order["status"] == "created"
                and order.get("expires_at", 0) > 0
                and order["expires_at"] <= now
            ):
                order["status"] = "expired"
                order["task_status"] = "none"
                changed = True
        if changed:
            self._save_locked()
        return changed

    def _item_by_id_locked(self, item_id: int) -> Dict[str, Any] | None:
        return next((item for item in self._state["items"] if item["id"] == item_id), None)

    def _order_by_paypal_id_locked(self, paypal_order_id: str) -> Dict[str, Any] | None:
        return next(
            (order for order in self._state["orders"] if order.get("paypal_order_id") == paypal_order_id),
            None,
        )

    def _order_counts_locked(self, order: Dict[str, Any]) -> bool:
        if order.get("is_test", False):
            return bool(self._state.get("test_mode", False))
        return True

    def _item_sales_locked(self, item_id: int) -> tuple[int, int, float]:
        now = time.time()
        paid = 0
        pending = 0
        last_paid_ts = 0.0
        for order in self._state["orders"]:
            if not self._order_counts_locked(order):
                continue
            if order["kind"] != "item" or order.get("item_id") != item_id:
                continue
            if order["status"] == "paid":
                paid += 1
                last_paid_ts = max(last_paid_ts, float(order.get("paid_ts", 0.0) or 0.0))
            elif order["status"] == "created" and float(order.get("expires_at", 0.0) or 0.0) > now:
                pending += 1
        return paid, pending, last_paid_ts

    def _event_state(self) -> Dict[str, Any]:
        if not self._get_event_state:
            return {}
        try:
            state = self._get_event_state()
            return state if isinstance(state, dict) else {}
        except Exception as exc:
            print(f"[SHOP] Event-State konnte nicht gelesen werden: {exc}")
            return {}

    def _public_item_locked(self, item: Dict[str, Any]) -> Dict[str, Any]:
        paid, pending, last_paid_ts = self._item_sales_locked(item["id"])
        now = time.time()
        max_qty = item["max_quantity"]
        remaining = None if max_qty <= 0 else max(0, max_qty - paid - pending)
        cooldown_remaining = 0
        if item["cooldown_seconds"] > 0 and last_paid_ts > 0:
            cooldown_remaining = max(
                0,
                int(last_paid_ts + item["cooldown_seconds"] - now),
            )

        available = bool(item["active"])
        reason = ""
        if not item["active"]:
            available = False
            reason = "Pausiert"
        elif max_qty > 0 and remaining <= 0:
            available = False
            reason = "Ausverkauft"
        elif cooldown_remaining > 0:
            available = False
            reason = "Cooldown"
        elif item["cooldown_seconds"] > 0 and pending > 0:
            available = False
            reason = "Reserviert"

        event_state = self._event_state()
        if (
            available
            and item.get("action_type") == "pause_minus"
            and bool(event_state.get("pause_active", False))
        ):
            available = False
            reason = "Pause aktiv – derzeit nicht verfügbar"

        return {
            "id": item["id"],
            "name": item["name"],
            "description": item["description"],
            "price_cents": item["price_cents"],
            "price": self._money(item["price_cents"]),
            "max_quantity": max_qty,
            "sold_count": paid,
            "pending_count": pending,
            "remaining": remaining,
            "cooldown_seconds": item["cooldown_seconds"],
            "cooldown_remaining": cooldown_remaining,
            "target": item["target"],
            "stream_text": item["stream_text"],
            "icon": item["icon"],
            "image_url": item.get("image_url", ""),
            "sort_order": int(item.get("sort_order", item["id"] * 10)),
            "action_type": item.get("action_type", ""),
            "action_value": int(item.get("action_value", 0) or 0),
            "active": item["active"],
            "available": available,
            "availability_reason": reason,
        }

    def _summary_locked(self) -> Dict[str, Any]:
        purchase_cents = 0
        donation_cents = 0
        for order in self._state["orders"]:
            if not self._order_counts_locked(order):
                continue
            if order["status"] != "paid":
                continue
            if order["kind"] == "item":
                purchase_cents += order["amount_cents"]
            else:
                donation_cents += order["amount_cents"]
        return {
            "purchase_cents": purchase_cents,
            "donation_cents": donation_cents,
            "paid_cents": purchase_cents + donation_cents,
            "purchase_total": self._money(purchase_cents),
            "donation_total": self._money(donation_cents),
            "paid_total": self._money(purchase_cents + donation_cents),
        }

    def _supporters_locked(self) -> list[Dict[str, Any]]:
        grouped: dict[str, Dict[str, Any]] = {}
        for order in self._state["orders"]:
            if not self._order_counts_locked(order):
                continue
            if order["status"] != "paid":
                continue
            display_name = (order.get("display_name") or "").strip()
            if not display_name:
                key = f"anon:{order['id']}"
                label = "Anonym"
            else:
                key = display_name.casefold()
                label = display_name
            current = grouped.get(key)
            if current is None:
                current = {
                    "name": label,
                    "amount_cents": 0,
                    "last_paid_ts": 0.0,
                }
                grouped[key] = current
            current["amount_cents"] += int(order["amount_cents"])
            current["last_paid_ts"] = max(
                float(current["last_paid_ts"]),
                float(order.get("paid_ts", 0.0) or 0.0),
            )
        supporters = list(grouped.values())
        supporters.sort(key=lambda item: item["last_paid_ts"], reverse=True)
        for supporter in supporters:
            supporter["amount"] = self._money(supporter["amount_cents"])
        return supporters[:150]

    def public_state(self) -> Dict[str, Any]:
        self._cleanup_expired_locked()
        items = [self._public_item_locked(item) for item in self._state["items"]]
        items.sort(key=lambda item: (int(item.get("sort_order", item["id"] * 10)), item["id"]))
        active_items = [item for item in items if item["active"]]

        public_orders: list[Dict[str, Any]] = []
        for order in reversed(self._state["orders"]):
            if not self._order_counts_locked(order):
                continue
            if order["status"] != "paid":
                continue
            public_orders.append({
                "id": order["id"],
                "kind": order["kind"],
                "item_id": order["item_id"],
                "item_name": order["item_name"],
                "target": order["target"],
                "amount_cents": order["amount_cents"],
                "amount": self._money(order["amount_cents"]),
                "display_name": order["display_name"] or "Anonym",
                "message": order["message"],
                "paid_at": order["paid_at"],
                "paid_ts": order["paid_ts"],
                "task_status": order["task_status"],
                "is_test": bool(order.get("is_test", False)),
                "action_type": order.get("action_type", ""),
                "action_value": int(order.get("action_value", 0) or 0),
                "action_status": order.get("action_status", "none"),
            })
            if len(public_orders) >= 100:
                break

        open_tasks = [
            deepcopy(order)
            for order in public_orders
            if order["kind"] == "item" and order["task_status"] in {"open", "in_progress"}
        ]
        return {
            "test_mode": bool(self._state.get("test_mode", False)),
            "paypal_configured": self.paypal_configured,
            "webhook_configured": self.webhook_configured,
            "paypal_mode": self.paypal_mode,
            "summary": self._summary_locked(),
            "supporters": self._supporters_locked(),
            "items": active_items,
            "items_admin": items,
            "orders": public_orders,
            "open_tasks": open_tasks,
            "open_task_count": len(open_tasks),
        }

    async def handle_control_message(self, message: Dict[str, Any]) -> tuple[bool, str]:
        msg_type = str(message.get("type", ""))
        try:
            if msg_type == "shop_item_save":
                await self._save_item(message)
                return True, ""
            if msg_type == "shop_item_toggle":
                await self._toggle_item(message)
                return True, ""
            if msg_type == "shop_item_delete":
                await self._delete_item(message)
                return True, ""
            if msg_type == "shop_task_status":
                await self._set_task_status(message)
                return True, ""
            if msg_type == "shop_test_mode":
                await self._set_test_mode(message)
                return True, ""
            if msg_type == "shop_test_reset":
                await self._reset_test_data()
                return True, ""
        except ShopError as exc:
            return False, str(exc)
        return False, "Unbekannte Shop-Aktion."

    async def _set_test_mode(self, message: Dict[str, Any]) -> None:
        enabled = bool(message.get("enabled", False))
        async with self._lock:
            self._state["test_mode"] = enabled
            self._save_locked()

    async def _reset_test_data(self) -> None:
        async with self._lock:
            self._state["orders"] = [
                order
                for order in self._state["orders"]
                if not order.get("is_test", False)
            ]
            max_order = max(
                (int(order["id"]) for order in self._state["orders"]),
                default=0,
            )
            self._state["next_order_id"] = max_order + 1
            self._save_locked()

    async def _save_item(self, message: Dict[str, Any]) -> None:
        name = str(message.get("name", "")).strip()
        if not name or len(name) > 80:
            raise ShopError("Artikelname fehlt oder ist zu lang.")
        description = str(message.get("description", "")).strip()[:320]
        stream_text = str(message.get("stream_text", "")).strip()[:220]
        icon = str(message.get("icon", "🎯")).strip()[:12] or "🎯"
        image_url = str(message.get("image_url", "")).strip()[:600]
        if image_url and not image_url.lower().startswith(("http://", "https://")):
            raise ShopError("Bild-URL muss mit http:// oder https:// beginnen.")

        target = str(message.get("target", "general"))
        if target not in {"tzmarty", "korsar", "both", "general", "on_site"}:
            raise ShopError("Ungültiges Ziel.")

        action_type = str(message.get("action_type", "")).strip().lower()
        if action_type not in {"", "pause_minus"}:
            raise ShopError("Ungültiger Aktionstyp.")

        price_cents = self._parse_money_to_cents(message.get("price"), 50, 100000)
        try:
            max_quantity = max(0, min(999, int(message.get("max_quantity", 0) or 0)))
            cooldown_minutes = max(0, min(1440, int(message.get("cooldown_minutes", 0) or 0)))
            sort_order = int(message.get("sort_order", 0) or 0)
            action_value = max(0, min(24 * 3600, int(message.get("action_value", 0) or 0)))
        except (TypeError, ValueError) as exc:
            raise ShopError("Anzahl/Cooldown/Sortierung/Aktionswert ist ungültig.") from exc

        if action_type == "pause_minus" and action_value <= 0:
            raise ShopError("Pausendieb benötigt einen Aktionswert in Sekunden.")

        active = bool(message.get("active", True))
        raw_id = message.get("id")

        async with self._lock:
            if raw_id in (None, "", 0, "0"):
                item_id = int(self._state["next_item_id"])
                self._state["next_item_id"] = item_id + 1
                now = self._now_text()
                self._state["items"].append({
                    "id": item_id,
                    "name": name,
                    "description": description,
                    "price_cents": price_cents,
                    "max_quantity": max_quantity,
                    "cooldown_seconds": cooldown_minutes * 60,
                    "target": target,
                    "stream_text": stream_text,
                    "icon": icon,
                    "image_url": image_url,
                    "sort_order": sort_order if sort_order else item_id * 10,
                    "action_type": action_type,
                    "action_value": action_value,
                    "active": active,
                    "created_at": now,
                    "updated_at": now,
                })
            else:
                try:
                    item_id = int(raw_id)
                except (TypeError, ValueError) as exc:
                    raise ShopError("Ungültige Artikel-ID.") from exc
                item = self._item_by_id_locked(item_id)
                if not item:
                    raise ShopError("Artikel nicht gefunden.")
                item.update({
                    "name": name,
                    "description": description,
                    "price_cents": price_cents,
                    "max_quantity": max_quantity,
                    "cooldown_seconds": cooldown_minutes * 60,
                    "target": target,
                    "stream_text": stream_text,
                    "icon": icon,
                    "image_url": image_url,
                    "sort_order": sort_order if sort_order else item_id * 10,
                    "action_type": action_type,
                    "action_value": action_value,
                    "active": active,
                    "updated_at": self._now_text(),
                })
            self._save_locked()

    async def _toggle_item(self, message: Dict[str, Any]) -> None:
        try:
            item_id = int(message.get("id", 0))
        except (TypeError, ValueError) as exc:
            raise ShopError("Ungültige Artikel-ID.") from exc
        async with self._lock:
            item = self._item_by_id_locked(item_id)
            if not item:
                raise ShopError("Artikel nicht gefunden.")
            item["active"] = not bool(item["active"])
            item["updated_at"] = self._now_text()
            self._save_locked()

    async def _delete_item(self, message: Dict[str, Any]) -> None:
        try:
            item_id = int(message.get("id", 0))
        except (TypeError, ValueError) as exc:
            raise ShopError("Ungültige Artikel-ID.") from exc
        async with self._lock:
            used = any(order.get("item_id") == item_id for order in self._state["orders"])
            if used:
                raise ShopError("Artikel wurde bereits verwendet. Bitte stattdessen deaktivieren.")
            before = len(self._state["items"])
            self._state["items"] = [item for item in self._state["items"] if item["id"] != item_id]
            if len(self._state["items"]) == before:
                raise ShopError("Artikel nicht gefunden.")
            self._save_locked()

    async def _set_task_status(self, message: Dict[str, Any]) -> None:
        try:
            order_id = int(message.get("order_id", 0))
        except (TypeError, ValueError) as exc:
            raise ShopError("Ungültige Kauf-ID.") from exc
        status = str(message.get("status", ""))
        if status not in {"open", "in_progress", "done"}:
            raise ShopError("Ungültiger Aufgabenstatus.")
        async with self._lock:
            order = next((o for o in self._state["orders"] if o["id"] == order_id), None)
            if not order or order["kind"] != "item" or order["status"] != "paid":
                raise ShopError("Aufgabe nicht gefunden.")
            order["task_status"] = status
            self._save_locked()

    def _public_base(self, request: web.Request) -> str:
        if self.public_base_url:
            return self.public_base_url
        forwarded = request.headers.get("X-Forwarded-Proto", "").strip()
        scheme = forwarded or request.scheme
        return f"{scheme}://{request.host}"

    async def _paypal_access_token(self) -> str:
        if not self.paypal_configured:
            raise ShopError("PayPal ist noch nicht konfiguriert.")
        now = time.time()
        if self._paypal_token and self._paypal_token_expires_at > now + 30:
            return self._paypal_token
        auth = aiohttp.BasicAuth(self.paypal_client_id, self.paypal_client_secret)
        timeout = aiohttp.ClientTimeout(total=20)
        async with aiohttp.ClientSession(timeout=timeout) as session:
            async with session.post(
                f"{self.paypal_api_base}/v1/oauth2/token",
                auth=auth,
                data={"grant_type": "client_credentials"},
                headers={"Accept": "application/json"},
            ) as response:
                data = await response.json(content_type=None)
                if response.status >= 300:
                    raise ShopError("PayPal-Zugriffstoken konnte nicht erstellt werden.")
        token = str(data.get("access_token", ""))
        if not token:
            raise ShopError("PayPal hat kein Zugriffstoken geliefert.")
        self._paypal_token = token
        self._paypal_token_expires_at = now + max(60, int(data.get("expires_in", 300) or 300))
        return token

    async def _paypal_request(
        self,
        method: str,
        path: str,
        payload: Dict[str, Any] | None = None,
        request_id: str | None = None,
    ) -> tuple[int, Dict[str, Any]]:
        token = await self._paypal_access_token()
        headers = {
            "Authorization": f"Bearer {token}",
            "Content-Type": "application/json",
            "Accept": "application/json",
        }
        if request_id:
            headers["PayPal-Request-Id"] = request_id[:108]
        timeout = aiohttp.ClientTimeout(total=25)
        async with aiohttp.ClientSession(timeout=timeout) as session:
            async with session.request(
                method,
                f"{self.paypal_api_base}{path}",
                headers=headers,
                json=payload,
            ) as response:
                try:
                    data = await response.json(content_type=None)
                except Exception:
                    data = {}
                return response.status, data

    async def _create_paypal_order(
        self,
        request: web.Request,
        local_order: Dict[str, Any],
    ) -> str:
        local_id = local_order["id"]
        base = self._public_base(request)
        return_url = f"{base}{self.prefix}/shop/paypal/return?{urlencode({'local_order': local_id})}"
        cancel_url = f"{base}{self.prefix}/shop/paypal/cancel?{urlencode({'local_order': local_id})}"
        if local_order["kind"] == "item":
            description = f"Foltershop: {local_order['item_name']}"
        else:
            description = "Spende für den Dartmarathon"
        payload = {
            "intent": "CAPTURE",
            "purchase_units": [{
                "reference_id": f"dm-{local_id}",
                "custom_id": f"dm-{local_id}",
                "description": description[:127],
                "amount": {
                    "currency_code": "EUR",
                    "value": self._money(local_order["amount_cents"]),
                },
            }],
            "application_context": {
                "brand_name": "Dartmarathon Foltershop",
                "landing_page": "LOGIN",
                "shipping_preference": "NO_SHIPPING",
                "user_action": "PAY_NOW",
                "return_url": return_url,
                "cancel_url": cancel_url,
            },
        }
        status, data = await self._paypal_request(
            "POST",
            "/v2/checkout/orders",
            payload,
            request_id=f"dm-create-{local_id}-{secrets.token_hex(6)}",
        )
        if status not in {200, 201}:
            print(f"[SHOP] PayPal create failed: {status} {data}")
            raise ShopError("PayPal konnte die Zahlung nicht starten.")
        paypal_id = str(data.get("id", ""))
        approval = ""
        for link in data.get("links", []) or []:
            if link.get("rel") in {"approve", "payer-action"}:
                approval = str(link.get("href", ""))
                break
        if not paypal_id or not approval:
            raise ShopError("PayPal hat keinen Freigabelink geliefert.")
        async with self._lock:
            current = next((o for o in self._state["orders"] if o["id"] == local_id), None)
            if not current:
                raise ShopError("Lokale Bestellung nicht gefunden.")
            current["paypal_order_id"] = paypal_id
            self._save_locked()
        return approval

    async def _new_local_order(
        self,
        kind: str,
        item_id: int,
        amount_cents: int,
        display_name: str,
        message: str,
    ) -> Dict[str, Any]:
        async with self._lock:
            self._cleanup_expired_locked()
            item_name = ""
            target = "general"
            stream_text = ""
            action_type = ""
            action_value = 0
            if kind == "item":
                item = self._item_by_id_locked(item_id)
                if not item:
                    raise ShopError("Artikel nicht gefunden.")
                public_item = self._public_item_locked(item)
                if not public_item["available"]:
                    reason = public_item["availability_reason"] or "aktuell nicht verfügbar"
                    raise ShopError(f"Artikel ist {reason.lower()}.")
                amount_cents = item["price_cents"]
                item_name = item["name"]
                target = item["target"]
                stream_text = item["stream_text"] or item["description"] or item["name"]
                action_type = item.get("action_type", "")
                action_value = int(item.get("action_value", 0) or 0)

            order_id = int(self._state["next_order_id"])
            self._state["next_order_id"] = order_id + 1
            now_ts = time.time()
            order = {
                "id": order_id,
                "kind": kind,
                "item_id": item_id if kind == "item" else 0,
                "item_name": item_name,
                "target": target,
                "stream_text": stream_text,
                "amount_cents": amount_cents,
                "display_name": display_name,
                "message": message,
                "status": "created",
                "task_status": "none",
                "created_at": self._now_text(),
                "created_ts": now_ts,
                "expires_at": now_ts + self.reservation_seconds,
                "paid_at": "",
                "paid_ts": 0.0,
                "paypal_order_id": "",
                "capture_id": "",
                "is_test": False,
                "action_type": action_type,
                "action_value": action_value,
                "action_status": "none",
            }
            self._state["orders"].append(order)
            self._save_locked()
            return deepcopy(order)

    async def _mark_failed(self, local_order_id: int) -> None:
        async with self._lock:
            order = next((o for o in self._state["orders"] if o["id"] == local_order_id), None)
            if order and order["status"] == "created":
                order["status"] = "failed"
                self._save_locked()

    async def _mark_cancelled(self, local_order_id: int) -> None:
        async with self._lock:
            order = next((o for o in self._state["orders"] if o["id"] == local_order_id), None)
            if order and order["status"] == "created":
                order["status"] = "cancelled"
                self._save_locked()

    async def _mark_paid(
        self,
        paypal_order_id: str,
        capture_id: str,
    ) -> Dict[str, Any] | None:
        order_public: Dict[str, Any] | None = None
        popup_payload: Dict[str, Any] | None = None
        async with self._lock:
            order = self._order_by_paypal_id_locked(paypal_order_id)
            if not order:
                print(f"[SHOP] PayPal order unbekannt: {paypal_order_id}")
                return None
            if order["status"] == "paid":
                return deepcopy(order)
            if capture_id and any(
                other.get("capture_id") == capture_id and other["status"] == "paid"
                for other in self._state["orders"]
                if other["id"] != order["id"]
            ):
                return None
            order["status"] = "paid"
            order["capture_id"] = capture_id
            order["paid_at"] = self._now_text()
            order["paid_ts"] = time.time()
            order["expires_at"] = 0.0
            order["task_status"] = "open" if order["kind"] == "item" else "none"
            if order["kind"] == "item" and order.get("action_type") == "pause_minus":
                order["action_status"] = "pending"
            else:
                order["action_status"] = "none"
            self._save_locked()
            order_public = deepcopy(order)
            name = order["display_name"] or "Anonym"
            if order["kind"] == "item":
                popup_payload = {
                    "type": "shop_purchase_popup",
                    "duration_ms": 20000,
                    "name": name,
                    "amount": self._money(order["amount_cents"]),
                    "item_name": order["item_name"],
                    "target": order["target"],
                    "stream_text": order["stream_text"] or order["item_name"],
                    "message": order["message"],
                }
            else:
                popup_payload = {
                    "type": "shop_donation_popup",
                    "duration_ms": 20000,
                    "name": name,
                    "amount": self._money(order["amount_cents"]),
                    "message": order["message"],
                }

        if order_public and order_public.get("action_status") == "pending":
            await self._try_apply_order_action(int(order_public["id"]))

        if self._broadcast_state:
            await self._broadcast_state()
        if popup_payload and self._broadcast_payload:
            await self._broadcast_payload(popup_payload)
        return order_public

    async def _try_apply_order_action(self, order_id: int) -> bool:
        async with self._lock:
            order = next(
                (
                    entry
                    for entry in self._state["orders"]
                    if int(entry["id"]) == int(order_id)
                ),
                None,
            )
            if (
                not order
                or order.get("status") != "paid"
                or order.get("action_status") != "pending"
            ):
                return False

            action_type = str(order.get("action_type", ""))
            action_value = int(order.get("action_value", 0) or 0)

        if not action_type:
            return False
        if not self._apply_event_action:
            return False

        try:
            result = await self._apply_event_action(
                action_type,
                action_value,
                int(order_id),
            )
        except Exception as exc:
            print(f"[SHOP] Event-Aktion fehlgeschlagen: {exc}")
            return False

        if not isinstance(result, dict) or not bool(result.get("applied", False)):
            return False

        async with self._lock:
            order = next(
                (
                    entry
                    for entry in self._state["orders"]
                    if int(entry["id"]) == int(order_id)
                ),
                None,
            )
            if not order or order.get("action_status") != "pending":
                return False

            order["action_status"] = "applied"
            # Automatisch ausgeführte Pausendiebe gelten als erledigt.
            if order.get("task_status") in {"open", "in_progress"}:
                order["task_status"] = "done"
            self._save_locked()

        return True

    async def process_pending_actions(self) -> None:
        async with self._lock:
            pending_ids = [
                int(order["id"])
                for order in self._state["orders"]
                if (
                    order.get("status") == "paid"
                    and order.get("action_status") == "pending"
                )
            ]

        for order_id in pending_ids[:20]:
            await self._try_apply_order_action(order_id)

    async def shop_page(self, request: web.Request) -> web.Response:
        path = self.template_dir / "shop.html"
        if not path.exists():
            raise web.HTTPNotFound()
        response = web.Response(text=path.read_text(encoding="utf-8"), content_type="text/html")
        response.headers["Cache-Control"] = "no-store"
        return response

    async def _complete_test_order(
        self,
        local_order_id: int,
    ) -> Dict[str, Any] | None:
        synthetic_order_id = (
            f"TEST-ORDER-{local_order_id}-"
            f"{secrets.token_hex(4).upper()}"
        )
        synthetic_capture_id = (
            f"TEST-CAPTURE-{local_order_id}-"
            f"{secrets.token_hex(4).upper()}"
        )

        async with self._lock:
            order = next(
                (
                    entry
                    for entry in self._state["orders"]
                    if entry["id"] == local_order_id
                ),
                None,
            )
            if not order:
                return None
            if order["status"] != "created":
                return deepcopy(order)
            order["paypal_order_id"] = synthetic_order_id
            order["is_test"] = True
            self._save_locked()

        return await self._mark_paid(
            synthetic_order_id,
            synthetic_capture_id,
        )

    async def create_order(self, request: web.Request) -> web.Response:
        test_mode = bool(self._state.get("test_mode", False))
        if not test_mode and not self.paypal_configured:
            return web.json_response(
                {
                    "ok": False,
                    "error": "PayPal ist noch nicht konfiguriert.",
                },
                status=503,
            )
        try:
            data = await request.json()
        except Exception:
            return web.json_response({"ok": False, "error": "Ungültige Anfrage."}, status=400)

        kind = str(data.get("kind", ""))
        display_name = str(data.get("display_name", "")).strip()[:50]
        message = str(data.get("message", "")).strip()[:180]
        try:
            if kind == "item":
                item_id = int(data.get("item_id", 0))
                amount_cents = 0
            elif kind == "donation":
                item_id = 0
                amount_cents = self._parse_money_to_cents(data.get("amount"), 100, 500000)
            else:
                raise ShopError("Ungültige Kaufart.")

            order = await self._new_local_order(
                kind,
                item_id,
                amount_cents,
                display_name,
                message,
            )

            if test_mode:
                completed = await self._complete_test_order(
                    int(order["id"])
                )
                if not completed:
                    raise ShopError(
                        "Testkauf konnte nicht abgeschlossen werden."
                    )
                return web.json_response(
                    {
                        "ok": True,
                        "test_completed": True,
                        "order_id": int(order["id"]),
                    }
                )

            try:
                approval_url = await self._create_paypal_order(request, order)
            except Exception:
                await self._mark_failed(order["id"])
                raise
            if self._broadcast_state:
                await self._broadcast_state()
            return web.json_response({"ok": True, "approval_url": approval_url})
        except ShopError as exc:
            return web.json_response({"ok": False, "error": str(exc)}, status=400)
        except Exception as exc:
            print(f"[SHOP] create_order error: {exc}")
            return web.json_response({"ok": False, "error": "Zahlung konnte nicht gestartet werden."}, status=500)

    async def _capture_and_process(self, paypal_order_id: str) -> bool:
        async with self._lock:
            local = self._order_by_paypal_id_locked(paypal_order_id)
            if not local:
                return False
            if local["status"] == "paid":
                return True
            local_id = int(local["id"])

        status, data = await self._paypal_request(
            "POST",
            f"/v2/checkout/orders/{paypal_order_id}/capture",
            {},
            request_id=f"dm-capture-{local_id}",
        )

        if status not in {200, 201}:
            # Ein paralleler Return/Webhook kann denselben Capture bereits abgeschlossen haben.
            async with self._lock:
                current = self._order_by_paypal_id_locked(paypal_order_id)
                if current and current["status"] == "paid":
                    return True
            print(f"[SHOP] capture failed: {status} {data}")
            return False

        capture_id = ""
        capture_status = ""
        for unit in data.get("purchase_units", []) or []:
            for capture in (((unit.get("payments") or {}).get("captures")) or []):
                capture_id = str(capture.get("id", "")) or capture_id
                capture_status = str(capture.get("status", "")) or capture_status
                if capture_status == "COMPLETED":
                    break

        if str(data.get("status", "")) == "COMPLETED" or capture_status == "COMPLETED":
            await self._mark_paid(paypal_order_id, capture_id)
            return True

        return False

    async def paypal_return(self, request: web.Request) -> web.StreamResponse:
        paypal_order_id = str(request.query.get("token", "")).strip()
        try:
            local_order_id = int(request.query.get("local_order", "0") or 0)
        except ValueError:
            local_order_id = 0
        if not paypal_order_id or local_order_id <= 0:
            raise web.HTTPFound(f"{self.prefix}/shop?payment=error")

        async with self._lock:
            local = next((o for o in self._state["orders"] if o["id"] == local_order_id), None)
            if not local or local.get("paypal_order_id") != paypal_order_id:
                raise web.HTTPFound(f"{self.prefix}/shop?payment=error")
            if local["status"] == "paid":
                raise web.HTTPFound(f"{self.prefix}/shop?payment=success")

        try:
            completed = await self._capture_and_process(paypal_order_id)
        except ShopError as exc:
            print(f"[SHOP] capture error: {exc}")
            raise web.HTTPFound(f"{self.prefix}/shop?payment=error")

        if completed:
            raise web.HTTPFound(f"{self.prefix}/shop?payment=success")
        raise web.HTTPFound(f"{self.prefix}/shop?payment=pending")

    async def paypal_cancel(self, request: web.Request) -> web.StreamResponse:
        try:
            local_order_id = int(request.query.get("local_order", "0") or 0)
        except ValueError:
            local_order_id = 0
        if local_order_id > 0:
            await self._mark_cancelled(local_order_id)
            if self._broadcast_state:
                await self._broadcast_state()
        raise web.HTTPFound(f"{self.prefix}/shop?payment=cancelled")

    async def _verify_webhook(self, request: web.Request, event: Dict[str, Any]) -> bool:
        if not self.webhook_configured:
            return False
        payload = {
            "auth_algo": request.headers.get("PAYPAL-AUTH-ALGO", ""),
            "cert_url": request.headers.get("PAYPAL-CERT-URL", ""),
            "transmission_id": request.headers.get("PAYPAL-TRANSMISSION-ID", ""),
            "transmission_sig": request.headers.get("PAYPAL-TRANSMISSION-SIG", ""),
            "transmission_time": request.headers.get("PAYPAL-TRANSMISSION-TIME", ""),
            "webhook_id": self.paypal_webhook_id,
            "webhook_event": event,
        }
        status, data = await self._paypal_request(
            "POST",
            "/v1/notifications/verify-webhook-signature",
            payload,
        )
        return status in {200, 201} and str(data.get("verification_status", "")) == "SUCCESS"

    async def paypal_webhook(self, request: web.Request) -> web.Response:
        try:
            event = await request.json()
        except Exception:
            return web.Response(status=400, text="invalid json")
        try:
            verified = await self._verify_webhook(request, event)
        except Exception as exc:
            print(f"[SHOP] webhook verify error: {exc}")
            return web.Response(status=503, text="verification unavailable")
        if not verified:
            return web.Response(status=400, text="invalid signature")

        event_type = str(event.get("event_type", ""))
        resource = event.get("resource") or {}

        if event_type == "CHECKOUT.ORDER.APPROVED":
            paypal_order_id = str(resource.get("id", ""))
            if paypal_order_id:
                try:
                    await self._capture_and_process(paypal_order_id)
                except Exception as exc:
                    print(f"[SHOP] approved webhook capture error: {exc}")
                    return web.Response(status=503, text="capture unavailable")

        elif event_type == "PAYMENT.CAPTURE.COMPLETED":
            capture_id = str(resource.get("id", ""))
            related = ((resource.get("supplementary_data") or {}).get("related_ids") or {})
            paypal_order_id = str(related.get("order_id", ""))
            if paypal_order_id:
                await self._mark_paid(paypal_order_id, capture_id)

        return web.Response(status=200, text="ok")

    @staticmethod
    def _normalize_header(value: Any) -> str:
        text = str(value or "").strip().casefold()
        text = re.sub(r"\s+", " ", text)
        text = text.replace("_", " ").replace("-", " ")
        text = re.sub(r"\s*\+\s*", " + ", text)
        return text.strip()

    @staticmethod
    def _parse_active(value: Any, default: bool = True) -> bool:
        text = str(value or "").strip().casefold()
        if not text:
            return default
        if text in {"ja", "j", "yes", "y", "true", "1", "aktiv"}:
            return True
        if text in {"nein", "n", "no", "false", "0", "inaktiv", "pausiert"}:
            return False
        raise ShopError(f"Aktiv-Wert '{value}' ist ungültig. Erlaubt: JA/NEIN.")

    @staticmethod
    def _parse_target(value: Any) -> str:
        text = str(value or "").strip().casefold()
        aliases = {
            "": "general",
            "allgemein": "general",
            "general": "general",
            "tzmarty": "tzmarty",
            "korsar": "korsar",
            "beide": "both",
            "both": "both",
            "vor ort": "on_site",
            "vorort": "on_site",
            "on site": "on_site",
            "on_site": "on_site",
        }
        if text not in aliases:
            raise ShopError(
                f"Ziel '{value}' ist ungültig. "
                "Erlaubt: Tzmarty, Korsar, Beide, Allgemein, Vor Ort."
            )
        return aliases[text]

    @staticmethod
    def _parse_action_type(value: Any) -> str:
        text = str(value or "").strip().casefold().replace("-", "_").replace(" ", "_")
        aliases = {
            "": "",
            "keine": "",
            "none": "",
            "pause_minus": "pause_minus",
            "pausendieb": "pause_minus",
        }
        if text not in aliases:
            raise ShopError(
                f"Aktionstyp '{value}' ist ungültig. Erlaubt: leer oder pause_minus."
            )
        return aliases[text]

    def _canonical_import_headers(self, headers: list[Any]) -> dict[str, int]:
        aliases = {
            "id": {"id", "artikel id", "artikelid"},
            "active": {"aktiv"},
            "max_quantity": {"anzahl", "maximalanzahl", "max anzahl", "maximum"},
            "name": {"name", "artikel", "artikelname"},
            "description": {"beschreibung"},
            "description_stream": {
                "beschreibung + streamtext",
                "beschreibung streamtext",
                "beschreibung+streamtext",
            },
            "stream_text": {"streamtext", "stream text"},
            "price": {"preis", "price"},
            "cooldown": {"cooldown", "cool down"},
            "target": {"ziel", "target"},
            "image_url": {"bild url", "bildurl", "image url", "imageurl", "bild"},
            "sort_order": {"sortierung", "reihenfolge", "sort order"},
            "action_type": {"aktionstyp", "aktions typ", "action type"},
            "action_value": {"aktionswert", "aktions wert", "action value"},
            "icon": {"icon", "emoji", "symbol"},
        }

        result: dict[str, int] = {}
        for index, raw in enumerate(headers):
            normalized = self._normalize_header(raw)
            if not normalized:
                continue
            for canonical, names in aliases.items():
                if normalized in names:
                    if canonical in result:
                        raise ShopError(
                            f"Spalte '{raw}' ist doppelt bzw. entspricht einer bereits erkannten Spalte."
                        )
                    result[canonical] = index
                    break

        required = ["id", "name", "price"]
        missing = [name for name in required if name not in result]
        if missing:
            labels = {"id": "ID", "name": "Name", "price": "Preis"}
            raise ShopError(
                "Pflichtspalte(n) fehlen: "
                + ", ".join(labels[name] for name in missing)
            )

        if "description" not in result and "description_stream" not in result:
            raise ShopError(
                "Es fehlt 'Beschreibung' oder 'Beschreibung + Streamtext'."
            )

        return result

    def _read_import_table(
        self,
        filename: str,
        payload: bytes,
    ) -> tuple[list[Any], list[list[Any]], dict[int, str]]:
        suffix = Path(filename or "").suffix.casefold()
        if suffix == ".csv":
            try:
                text = payload.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = payload.decode("cp1252")

            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\\t,")
            except csv.Error:
                dialect = csv.excel
                dialect.delimiter = ";"

            reader = csv.reader(io.StringIO(text), dialect)
            rows = list(reader)
            if not rows:
                raise ShopError("Die CSV-Datei ist leer.")
            return list(rows[0]), [list(row) for row in rows[1:]], {}

        if suffix != ".xlsx":
            raise ShopError("Erlaubt sind nur .xlsx und .csv.")

        try:
            from openpyxl import load_workbook
        except Exception as exc:
            raise ShopError(
                "XLSX-Import ist auf dem Server noch nicht verfügbar. "
                "Bitte requirements.txt mit openpyxl deployen."
            ) from exc

        try:
            workbook = load_workbook(
                io.BytesIO(payload),
                data_only=True,
                read_only=False,
            )
        except Exception as exc:
            raise ShopError("Die XLSX-Datei konnte nicht gelesen werden.") from exc

        sheet = workbook.active
        raw_rows = list(sheet.iter_rows())
        if not raw_rows:
            raise ShopError("Die Excel-Datei ist leer.")

        headers = [cell.value for cell in raw_rows[0]]
        rows: list[list[Any]] = []
        image_links: dict[int, str] = {}

        header_map = self._canonical_import_headers(headers)
        image_index = header_map.get("image_url")

        for excel_row_number, cells in enumerate(raw_rows[1:], start=2):
            values = [cell.value for cell in cells]
            if not any(value not in (None, "") for value in values):
                continue
            rows.append(values)
            if (
                image_index is not None
                and image_index < len(cells)
                and getattr(cells[image_index], "hyperlink", None)
            ):
                target = str(cells[image_index].hyperlink.target or "").strip()
                if target:
                    image_links[excel_row_number] = target

        return headers, rows, image_links

    def _normalize_import_rows(
        self,
        headers: list[Any],
        rows: list[list[Any]],
        image_links: dict[int, str] | None = None,
    ) -> tuple[list[Dict[str, Any]], list[str], list[str]]:
        mapping = self._canonical_import_headers(headers)
        image_links = image_links or {}
        normalized: list[Dict[str, Any]] = []
        errors: list[str] = []
        warnings: list[str] = []
        seen_ids: set[int] = set()

        def value(row: list[Any], key: str, default: Any = "") -> Any:
            index = mapping.get(key)
            if index is None or index >= len(row):
                return default
            return row[index]

        for offset, row in enumerate(rows, start=2):
            if not any(cell not in (None, "") for cell in row):
                continue
            if len(normalized) >= 500:
                errors.append("Es können maximal 500 Artikel pro Datei importiert werden.")
                break

            try:
                raw_id = value(row, "id")
                item_id = int(float(str(raw_id).replace(",", ".")))
                if item_id <= 0:
                    raise ShopError("ID muss größer als 0 sein.")
                if item_id in seen_ids:
                    raise ShopError(f"ID {item_id} kommt in der Datei mehrfach vor.")
                seen_ids.add(item_id)

                name = str(value(row, "name")).strip()
                if not name or len(name) > 80:
                    raise ShopError("Name fehlt oder ist länger als 80 Zeichen.")

                combined = str(value(row, "description_stream")).strip()
                description = str(value(row, "description")).strip() or combined
                if not description:
                    raise ShopError("Beschreibung fehlt.")
                description = description[:320]

                stream_text = str(value(row, "stream_text")).strip()
                if not stream_text:
                    stream_text = combined or description
                stream_text = stream_text[:220]

                price_cents = self._parse_money_to_cents(
                    value(row, "price"),
                    50,
                    100000,
                )

                try:
                    max_quantity = max(
                        0,
                        min(
                            999,
                            int(float(str(value(row, "max_quantity", 0) or 0).replace(",", "."))),
                        ),
                    )
                    cooldown_minutes = max(
                        0,
                        min(
                            1440,
                            int(float(str(value(row, "cooldown", 0) or 0).replace(",", "."))),
                        ),
                    )
                except (TypeError, ValueError) as exc:
                    raise ShopError("Anzahl oder CoolDown ist keine gültige Zahl.") from exc

                active = self._parse_active(value(row, "active"), True)
                target = self._parse_target(value(row, "target"))
                action_type = self._parse_action_type(value(row, "action_type"))

                try:
                    raw_action_value = value(row, "action_value", 0)
                    action_value = max(
                        0,
                        min(
                            24 * 3600,
                            int(float(str(raw_action_value or 0).replace(",", "."))),
                        ),
                    )
                except (TypeError, ValueError) as exc:
                    raise ShopError("Aktionswert ist keine gültige Zahl.") from exc

                if action_type == "pause_minus" and action_value <= 0:
                    raise ShopError(
                        "Aktionstyp pause_minus benötigt Aktionswert in Sekunden."
                    )

                try:
                    raw_sort = value(row, "sort_order", item_id * 10)
                    sort_order = int(float(str(raw_sort or item_id * 10).replace(",", ".")))
                except (TypeError, ValueError) as exc:
                    raise ShopError("Sortierung ist keine gültige Zahl.") from exc

                image_url = str(
                    image_links.get(offset)
                    or value(row, "image_url")
                    or ""
                ).strip()[:600]
                if image_url and not image_url.lower().startswith(("http://", "https://")):
                    raise ShopError("Bild-Url muss mit http:// oder https:// beginnen.")

                icon = str(value(row, "icon", "🎯") or "🎯").strip()[:12] or "🎯"

                if (
                    not action_type
                    and name.casefold().startswith("pausendieb")
                ):
                    warnings.append(
                        f"Zeile {offset}: '{name}' ist ein Pausendieb, "
                        "hat aber keinen Aktionstyp. "
                        "Für automatische Pausenwirkung: Aktionstyp=pause_minus "
                        "und Aktionswert in Sekunden setzen."
                    )

                normalized.append({
                    "id": item_id,
                    "name": name,
                    "description": description,
                    "price_cents": price_cents,
                    "max_quantity": max_quantity,
                    "cooldown_seconds": cooldown_minutes * 60,
                    "target": target,
                    "stream_text": stream_text,
                    "icon": icon,
                    "image_url": image_url,
                    "sort_order": sort_order,
                    "action_type": action_type,
                    "action_value": action_value,
                    "active": active,
                })

            except ShopError as exc:
                errors.append(f"Zeile {offset}: {exc}")
            except Exception as exc:
                errors.append(f"Zeile {offset}: Ungültige Daten ({exc}).")

        return normalized, errors, warnings

    def _import_preview_locked(
        self,
        normalized: list[Dict[str, Any]],
    ) -> Dict[str, Any]:
        existing = {
            int(item["id"]): item
            for item in self._state["items"]
        }
        new_count = 0
        updated_count = 0
        unchanged_count = 0
        deactivated_count = 0
        rows_preview: list[Dict[str, Any]] = []

        compare_keys = (
            "name",
            "description",
            "price_cents",
            "max_quantity",
            "cooldown_seconds",
            "target",
            "stream_text",
            "icon",
            "image_url",
            "sort_order",
            "action_type",
            "action_value",
            "active",
        )

        for incoming in normalized:
            current = existing.get(int(incoming["id"]))
            if current is None:
                status = "neu"
                new_count += 1
            else:
                changed = any(
                    current.get(key) != incoming.get(key)
                    for key in compare_keys
                )
                status = "aktualisiert" if changed else "unverändert"
                if changed:
                    updated_count += 1
                else:
                    unchanged_count += 1

                if bool(current.get("active", True)) and not bool(incoming["active"]):
                    deactivated_count += 1

            if len(rows_preview) < 50:
                rows_preview.append({
                    "id": int(incoming["id"]),
                    "name": incoming["name"],
                    "status": status,
                    "price_cents": int(incoming["price_cents"]),
                    "active": bool(incoming["active"]),
                    "action_type": incoming.get("action_type", ""),
                    "action_value": int(incoming.get("action_value", 0) or 0),
                })

        return {
            "total": len(normalized),
            "new": new_count,
            "updated": updated_count,
            "unchanged": unchanged_count,
            "deactivated": deactivated_count,
            "rows": rows_preview,
        }

    async def import_items(self, request: web.Request) -> web.Response:
        try:
            reader = await request.multipart()
        except Exception:
            return web.json_response(
                {"ok": False, "errors": ["Ungültiger Datei-Upload."]},
                status=400,
            )

        action = "preview"
        filename = ""
        file_bytes = b""

        try:
            while True:
                part = await reader.next()
                if part is None:
                    break
                if part.name == "action":
                    action = (await part.text()).strip().lower() or "preview"
                elif part.name == "file":
                    filename = str(part.filename or "")
                    chunks = []
                    total = 0
                    while True:
                        chunk = await part.read_chunk(size=64 * 1024)
                        if not chunk:
                            break
                        total += len(chunk)
                        if total > 5 * 1024 * 1024:
                            raise ShopError("Importdatei darf maximal 5 MB groß sein.")
                        chunks.append(chunk)
                    file_bytes = b"".join(chunks)
        except ShopError as exc:
            return web.json_response(
                {"ok": False, "errors": [str(exc)]},
                status=400,
            )

        if action not in {"preview", "apply"}:
            return web.json_response(
                {"ok": False, "errors": ["Ungültige Importaktion."]},
                status=400,
            )
        if not filename or not file_bytes:
            return web.json_response(
                {"ok": False, "errors": ["Keine Datei ausgewählt."]},
                status=400,
            )

        try:
            headers, rows, image_links = self._read_import_table(filename, file_bytes)
            normalized, errors, warnings = self._normalize_import_rows(
                headers,
                rows,
                image_links,
            )
        except ShopError as exc:
            return web.json_response(
                {"ok": False, "errors": [str(exc)], "warnings": []},
                status=400,
            )

        if errors:
            return web.json_response(
                {
                    "ok": False,
                    "errors": errors[:100],
                    "warnings": warnings[:100],
                },
                status=400,
            )

        async with self._lock:
            preview = self._import_preview_locked(normalized)

            if action == "apply":
                now = self._now_text()
                by_id = {
                    int(item["id"]): item
                    for item in self._state["items"]
                }

                for incoming in normalized:
                    item_id = int(incoming["id"])
                    current = by_id.get(item_id)
                    if current is None:
                        current = {
                            **incoming,
                            "created_at": now,
                            "updated_at": now,
                        }
                        self._state["items"].append(current)
                        by_id[item_id] = current
                    else:
                        created_at = current.get("created_at", "")
                        current.clear()
                        current.update({
                            **incoming,
                            "created_at": created_at or now,
                            "updated_at": now,
                        })

                self._state["items"].sort(
                    key=lambda item: (
                        int(item.get("sort_order", item["id"] * 10)),
                        int(item["id"]),
                    )
                )
                max_id = max(
                    (int(item["id"]) for item in self._state["items"]),
                    default=0,
                )
                self._state["next_item_id"] = max(
                    int(self._state.get("next_item_id", 1)),
                    max_id + 1,
                )
                self._save_locked()

        if action == "apply" and self._broadcast_state:
            await self._broadcast_state()

        return web.json_response({
            "ok": True,
            "action": action,
            "filename": filename,
            "preview": preview,
            "warnings": warnings[:100],
        })

    def register_routes(self, app: web.Application) -> None:
        app.router.add_get(f"{self.prefix}/shop", self.shop_page)
        app.router.add_post(f"{self.prefix}/shop/api/create-order", self.create_order)
        app.router.add_get(f"{self.prefix}/shop/paypal/return", self.paypal_return)
        app.router.add_get(f"{self.prefix}/shop/paypal/cancel", self.paypal_cancel)
        app.router.add_post(f"{self.prefix}/paypal/webhook", self.paypal_webhook)
